from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.contrib import messages
import openpyxl
from .models import Client
from .forms import ClientForm
from django.shortcuts import redirect, get_object_or_404
from django.contrib import messages
from django.urls import reverse
from .models import Client

# ---------------- Client List ----------------
@login_required
def client_list(request):
    query = request.GET.get('q')
    clients = Client.objects.filter(is_deleted=False)  # Exclude soft-deleted clients

    if query:
        clients = clients.filter(
            Q(main_applicant__icontains=query) |
            Q(company__icontains=query)
        )

    # Ensure all clients have a status
    for client in clients:
        if not client.status:
            client.status = "Pending"

    return render(request, 'client_list.html', {
        'clients': clients,
    })



# ---------------- Create / Edit Client ----------------
@login_required
def client_form(request, pk=None):
    client = get_object_or_404(Client, pk=pk) if pk else None

    if request.method == 'POST':
        form = ClientForm(request.POST, instance=client)
        if form.is_valid():
            client_obj = form.save(commit=False)

            # Ensure optional fields are None if empty
            optional_fields = [
                'other_company', 'status', 'local_agent', 'gcbl_no',
                'cbi_submission_date', 'approval_date', 'fm', 'status_notes'
            ]
            for field in optional_fields:
                if not form.cleaned_data.get(field):
                    setattr(client_obj, field, None)

            client_obj.save()
            messages.success(request, f"Client {client_obj.main_applicant} saved successfully.")
            return redirect('client_list')
    else:
        # Pre-populate company and status from GET parameters if provided
        initial = {}
        company_param = request.GET.get('company')
        if company_param:
            initial['company'] = company_param
        
        status_param = request.GET.get('status')
        if status_param:
            initial['status'] = status_param
        
        form = ClientForm(instance=client, initial=initial)

    return render(request, 'client_form.html', {'form': form})


# ---------------- Company Views ----------------
@login_required
def company_list(request):
    # Extract choices from the model
    companies = [choice[0] for choice in Client.COMPANY_CHOICES]
    return render(request, 'company_list.html', {'companies': companies})

@login_required
def company_detail(request, company_name):
    query = request.GET.get('q')
    clients = Client.objects.filter(company=company_name, is_deleted=False)
    
    if query:
        clients = clients.filter(main_applicant__icontains=query)

    return render(request, 'company_detail.html', {
        'clients': clients,
        'company_name': company_name,
    })


# ---------------- Safe Delete ----------------
@login_required
def client_delete(request, pk):
    client = get_object_or_404(Client, pk=pk)
    client.is_deleted = True  # Soft delete
    client.save()
    messages.success(request, f"Client {client.main_applicant} deleted.")
    return redirect('client_list')


# ---------------- Restore Client ----------------
@login_required
def restore_client(request, pk):
    client = get_object_or_404(Client, pk=pk)
    client.is_deleted = False
    client.save()
    messages.success(request, f"Client {client.main_applicant} restored.")
    return redirect('client_list')


# ---------------- Export to Excel ----------------
def append_client_data(ws, clients):
    headers = [
        "ID", "Name", "Company", "Status", "Initial Payment",
        "Investment", "Local Agent", "GCBL No", "CBI Submission Date",
        "Approval Date", "FM", "Status Notes", "Created At"
    ]
    ws.append(headers)

    for client in clients:
        ws.append([
            client.id,
            client.main_applicant,
            client.company if client.company != 'Others' else client.other_company,
            client.status,
            client.initial_payment,
            client.investment,
            client.local_agent or '',
            client.gcbl_no or '',
            client.cbi_submission_date.strftime("%Y-%m-%d") if client.cbi_submission_date else '',
            client.approval_date.strftime("%Y-%m-%d") if client.approval_date else '',
            client.fm or '',
            client.status_notes or '',
            client.created_at.strftime("%Y-%m-%d %H:%M"),
        ])

@login_required
def export_excel(request):
    wb = openpyxl.Workbook()
    
    # 1. Status Sheets
    ws_approved = wb.active
    ws_approved.title = "Approved"
    append_client_data(ws_approved, Client.objects.filter(status='Approved', is_deleted=False))

    ws_issued = wb.create_sheet(title="Passport Issued")
    append_client_data(ws_issued, Client.objects.filter(status='Issued', is_deleted=False))

    # 3. Company Sheets
    for choice_val, choice_label in Client.COMPANY_CHOICES:
        # Excel sheet titles must be <= 31 chars and not contain certain chars
        sheet_title = choice_label[:31]
        ws_company = wb.create_sheet(title=sheet_title)
        append_client_data(ws_company, Client.objects.filter(company=choice_val, is_deleted=False))

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename=clients_report.xlsx'
    wb.save(response)
    return response


# ---------------- Transfer / Update Client Status ----------------

@login_required
def client_action(request, pk):
    client = get_object_or_404(Client, pk=pk)

    # Get redirect page from POST or default to client_list
    redirect_page = request.POST.get('redirect', 'client_list')

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'edit':
            return redirect('client_edit', pk=client.pk)
        elif action == 'delete':
            client.is_deleted = True
            client.save()
            messages.success(request, f"Client {client.main_applicant} deleted.")
        elif action == 'manual':
            # Only allow manual status update on client_list page
            if redirect_page == 'client_list':
                new_status = request.POST.get('manual_status')
                if new_status:
                    client.status = new_status
                    client.save()
                    messages.success(request, f"Status updated to {new_status}.")
        elif action == 'approve':
            client.status = 'Approved'
            client.save()
            messages.success(request, "Status changed to Approved.")
        elif action == 'issued':
            client.status = 'Issued'
            client.save()
            messages.success(request, "Status changed to Issued.")
        elif action == 'deny':
            client.status = 'Denied'
            client.save()
            messages.success(request, "Status changed to Denied.")

    # Redirect back to the same page
    if redirect_page == 'approved':
        return redirect('approved_clients')
    elif redirect_page == 'issued':
        return redirect('issued_clients')
    elif redirect_page == 'denied':
        return redirect('denied_clients')
    else:
        return redirect('client_list')



# ---------------- Pages for transferred clients ----------------
@login_required
def approved_clients(request):
    query = request.GET.get('q')
    clients = Client.objects.filter(status='Approved', is_deleted=False)
    if query:
        clients = clients.filter(
            Q(main_applicant__icontains=query) |
            Q(company__icontains=query)
        )
    return render(request, 'approved_clients.html', {'clients': clients})

@login_required
def issued_clients(request):
    query = request.GET.get('q')
    clients = Client.objects.filter(status='Issued', is_deleted=False)
    if query:
        clients = clients.filter(
            Q(main_applicant__icontains=query) |
            Q(company__icontains=query)
        )
    return render(request, 'issued_clients.html', {'clients': clients})

@login_required
def denied_clients(request):
    query = request.GET.get('q')
    clients = Client.objects.filter(status='Denied', is_deleted=False)
    if query:
        clients = clients.filter(
            Q(main_applicant__icontains=query) |
            Q(company__icontains=query)
        )
    return render(request, 'denied_clients.html', {'clients': clients})
